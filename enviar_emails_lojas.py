"""
Envia o relatorio de reabastecimento por e-mail, um e-mail por loja, via SMTP
(Outlook/Microsoft 365 por padrao).

NADA de senha fica gravado neste arquivo nem em lugar nenhum: o script pede a
senha (ou senha de app) na hora de rodar, na propria tela, e usa só naquele
momento pra autenticar no servidor.

Uso basico (primeiro sempre roda em --dry-run pra conferir antes de mandar de
verdade):

    python enviar_emails_lojas.py --onebeat "export_onebeat.xlsx" \
                                   --ritmo "RITMO POSICAO.csv" \
                                   --emails "emails_lojas.csv" \
                                   --remetente "compras@abra.com.br" \
                                   --dry-run

Quando a previa estiver ok, roda de novo sem --dry-run pra enviar de verdade:

    python enviar_emails_lojas.py --onebeat "export_onebeat.xlsx" \
                                   --ritmo "RITMO POSICAO.csv" \
                                   --emails "emails_lojas.csv" \
                                   --remetente "compras@abra.com.br"

O arquivo --emails e um CSV simples com duas colunas, uma linha por loja:

    loja,email
    ACV,loja.gaia@abra.com.br
    ACR,loja.dadi@abra.com.br
    ...

PARA ADICIONAR UMA LOJA NOVA: uma linha em STORE_SHEETS abaixo (igual no
gerar_reabastecimento.py) + uma linha no CSV de e-mails.
"""
import argparse
import getpass
import smtplib
import sys
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import openpyxl
import pandas as pd

# ---------------------------------------------------------------------------
STORE_SHEETS = ["ACV", "ACR", "MCO", "CSH", "REC", "ABI", "VIA", "ICA", "RIO", "SHT", "BSH"]
STORE_NAMES = {
    "ACV": "Gaia", "ACR": "Dadi", "MCO": "Eco", "CSH": "Bada", "REC": "Trip",
    "ABI": "Papel", "VIA": "Flip", "ICA": "Noga", "RIO": "Zoon", "SHT": "Arte", "BSH": "Zeni",
}

SMTP_PRESETS = {
    "office365": ("smtp.office365.com", 587),
    "gmail": ("smtp.gmail.com", 587),
}


def load_onebeat_export(path: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else h for h in rows[0]]
    df = pd.DataFrame(rows[1:], columns=header)
    rename = {
        "Loja": "Loja", "Codigo": "Codigo", "Código": "Codigo",
        "Rótulo": "Rotulo", "Rotulo": "Rotulo",
        "Alvo": "Alvo", "Estoque Loja": "EstoqueLoja", "Estoque CD": "EstoqueCD",
        "Reposição": "Reposicao", "Reposicao": "Reposicao",
    }
    df = df.rename(columns={c: rename.get(c, c) for c in df.columns})
    required = ["Loja", "Codigo", "Rotulo", "Reposicao"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        sys.exit(f"Export do onebeat sem as colunas esperadas: {missing}. Encontradas: {list(df.columns)}")
    df = df[df["Codigo"].notna() & df["Loja"].notna()]
    df = df[df["Loja"].astype(str).str.strip().str.lower() != "total"]
    df["Codigo"] = df["Codigo"].astype(str).str.strip()
    df["Reposicao"] = pd.to_numeric(df["Reposicao"], errors="coerce").fillna(0)
    df = df[df["Reposicao"] > 0]
    df = df.drop_duplicates(subset=["Loja", "Codigo"], keep="first")
    return df.reset_index(drop=True)


def load_ritmo_classes(path: str | None) -> dict:
    """Le codigo -> classe de um arquivo com a aba/CSV RITMO POSIÇÃO. Opcional."""
    if not path:
        return {}
    if path.lower().endswith(".csv"):
        raw = open(path, "rb").read()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        import csv, io
        delim = ";" if text.split("\n", 1)[0].count(";") > text.split("\n", 1)[0].count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delim)
        rows = list(reader)
    else:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = "RITMO POSIÇÃO" if "RITMO POSIÇÃO" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]

    header = [str(h or "").strip().lower() for h in rows[0]]
    def find(name):
        for i, h in enumerate(header):
            if h == name:
                return i
        return -1
    def find_contains(needle):
        for i, h in enumerate(header):
            if needle in h:
                return i
        return -1

    idx_codigo = find("recurso")
    if idx_codigo == -1:
        idx_codigo = find_contains("codigo")
    idx_classe = find_contains("classe")
    if idx_codigo == -1 or idx_classe == -1:
        return {}

    out = {}
    for row in rows[1:]:
        if not row or idx_codigo >= len(row):
            continue
        codigo = row[idx_codigo]
        if not codigo:
            continue
        codigo = str(codigo).strip()
        if codigo not in out:
            out[codigo] = row[idx_classe] if idx_classe < len(row) else ""
    return out


def load_emails(path: str) -> dict:
    df = pd.read_csv(path, dtype=str)
    df.columns = [c.strip().lower() for c in df.columns]
    if "loja" not in df.columns or "email" not in df.columns:
        sys.exit("O arquivo de e-mails precisa ter as colunas 'loja' e 'email'.")
    return {str(r["loja"]).strip().upper(): str(r["email"]).strip() for _, r in df.iterrows()}


def assign_store(loja_text: str) -> str | None:
    if not isinstance(loja_text, str):
        return None
    for token in loja_text.strip().split():
        if token.upper() in STORE_SHEETS:
            return token.upper()
    return None


def build_html_table(rows: pd.DataFrame, classes: dict) -> str:
    th = 'style="border:1px solid #999;background:#f0ead9;padding:6px 10px;text-align:left;font-family:Arial,sans-serif;font-size:13px;"'
    td = 'style="border:1px solid #999;padding:5px 10px;font-family:Arial,sans-serif;font-size:13px;"'
    td_num = 'style="border:1px solid #999;padding:5px 10px;font-family:Arial,sans-serif;font-size:13px;text-align:right;"'
    parts = ['<table style="border-collapse:collapse;">']
    parts.append("<tr>" + "".join(f"<th {th}>{h}</th>" for h in ["Código", "Nome", "Classe", "Enviar"]) + "</tr>")
    for _, r in rows.iterrows():
        classe = classes.get(r["Codigo"], "")
        parts.append(
            "<tr>"
            f"<td {td}>{r['Codigo']}</td>"
            f"<td {td}>{r['Rotulo']}</td>"
            f"<td {td}>{classe or ''}</td>"
            f"<td {td_num}>{int(r['Reposicao'])}</td>"
            "</tr>"
        )
    parts.append("</table>")
    return "".join(parts)


def send_email(smtp_host, smtp_port, remetente, senha, destinatario, assunto, html_body):
    msg = MIMEMultipart("alternative")
    msg["Subject"] = assunto
    msg["From"] = remetente
    msg["To"] = destinatario
    msg.attach(MIMEText("Este e-mail contém uma tabela HTML; abra em um cliente que exiba HTML.", "plain"))
    msg.attach(MIMEText(html_body, "html"))
    with smtplib.SMTP(smtp_host, smtp_port) as server:
        server.starttls()
        server.login(remetente, senha)
        server.sendmail(remetente, [destinatario], msg.as_string())


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--onebeat", required=True, help="export bruto do OneBeat (.xlsx)")
    ap.add_argument("--ritmo", help="arquivo/CSV com a aba RITMO POSIÇÃO (pra trazer a Classe)")
    ap.add_argument("--emails", required=True, help="CSV com colunas loja,email")
    ap.add_argument("--remetente", required=True, help="e-mail que vai disparar (a conta que vai logar no SMTP)")
    ap.add_argument("--servidor", choices=list(SMTP_PRESETS.keys()), default="office365", help="preset de servidor SMTP")
    ap.add_argument("--smtp-host", help="host SMTP manual (sobrepõe --servidor)")
    ap.add_argument("--smtp-port", type=int, help="porta SMTP manual")
    ap.add_argument("--dry-run", action="store_true", help="só mostra o que seria enviado, não manda nada")
    args = ap.parse_args()

    smtp_host, smtp_port = SMTP_PRESETS[args.servidor]
    if args.smtp_host:
        smtp_host = args.smtp_host
    if args.smtp_port:
        smtp_port = args.smtp_port

    print("Lendo export do OneBeat...")
    df = load_onebeat_export(args.onebeat)
    df["aba"] = df["Loja"].apply(assign_store)

    print("Lendo classes do RITMO POSIÇÃO (se fornecido)...")
    classes = load_ritmo_classes(args.ritmo)

    print("Lendo lista de e-mails das lojas...")
    emails = load_emails(args.emails)

    hoje = datetime.now().strftime("%d/%m/%Y")

    plano = []
    for store in STORE_SHEETS:
        sub = df[df["aba"] == store]
        email = emails.get(store)
        plano.append((store, STORE_NAMES.get(store, ""), len(sub), email, sub))

    print()
    print("=== PRÉVIA DO ENVIO ===")
    for store, nome, n, email, _ in plano:
        status = "OK" if n > 0 and email else ("SEM CÓDIGOS -- não será enviado" if n == 0 else "SEM E-MAIL CADASTRADO -- não será enviado")
        print(f"  {store} ({nome}): {n} código(s) -> {email or '(sem e-mail)'}  [{status}]")

    a_enviar = [(s, n, cnt, e, sub) for s, n, cnt, e, sub in plano if cnt > 0 and e]
    if not a_enviar:
        print("\nNenhuma loja com código + e-mail cadastrado. Nada a enviar.")
        return

    print(f"\n{len(a_enviar)} e-mail(s) seriam enviados.")
    if args.dry_run:
        print("Modo --dry-run: nada foi enviado.")
        return

    confirm = input(f"\nDigite ENVIAR para mandar esses {len(a_enviar)} e-mail(s) de verdade: ")
    if confirm.strip().upper() != "ENVIAR":
        print("Cancelado.")
        return

    senha = getpass.getpass(f"Senha (ou senha de app) de {args.remetente}: ")

    for store, nome, n, email, sub in a_enviar:
        assunto = f"Reabastecimento {store} ({nome}) - {hoje}"
        html = f"<p>Segue a sugestão de reabastecimento de hoje ({hoje}) para a loja {store} ({nome}):</p>" + build_html_table(sub, classes)
        try:
            send_email(smtp_host, smtp_port, args.remetente, senha, email, assunto, html)
            print(f"  [OK] {store} -> {email}")
        except Exception as exc:
            print(f"  [FALHOU] {store} -> {email}: {exc}")

    print("\nConcluído.")


if __name__ == "__main__":
    main()
